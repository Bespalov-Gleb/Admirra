"""Bounded owner/project-scoped quality report; detached before file rendering."""
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace
import sqlalchemy as sa
from core import models
from lead_validator.services.scoped_stats import _scope, rejected_condition
from lead_validator.services.scoped_placements import get_blacklist


def snapshot(db, owner_id, days):
    if isinstance(days, bool) or not isinstance(days, int) or not 1 <= days <= 365:
        raise ValueError('Invalid quality report period')
    if db.new or db.dirty or db.deleted:
        raise ValueError('Quality report requires a read-only workflow session')
    # The route has only authenticated the caller. End that read transaction,
    # capture all sections under one snapshot, then return only plain data.
    db.rollback()
    try:
        db.connection(execution_options={'isolation_level': 'REPEATABLE READ'})
        db.execute(sa.text('SET TRANSACTION READ ONLY'))
        end = datetime.now(timezone.utc)
        start = end - timedelta(days=days)
        lead, project = models.Lead, models.PhoneProject
        filters = _scope(owner_id, start, end)
        count = sa.func.count(lead.id)
        bad = count.filter(rejected_condition())
        total, rejected = db.execute(sa.select(count, bad).select_from(lead)
            .join(project, lead.project_id == project.id).where(*filters)).one()
        dims = [sa.func.coalesce(sa.func.nullif(col, ''), fallback) for col, fallback in
                ((lead.utm_source, 'direct'), (lead.utm_campaign, 'none'), (lead.utm_content, 'none'))]
        rows = db.execute(sa.select(project.id, sa.func.left(project.name, 200),
            *[sa.func.left(d, 200) for d in dims], count, bad).select_from(lead)
            .join(project, lead.project_id == project.id).where(*filters)
            .group_by(project.id, project.name, *dims).having(count >= 5, bad * 100 >= count * 50)
            .order_by((bad * 100.0 / count).desc(), bad.desc(), project.id, *dims).limit(10)).all()
        sources = [SimpleNamespace(project_id=str(pid), project_name=name, source=source, campaign=campaign,
            content=content, total_leads=amount, rejected_leads=bad_amount,
            rejection_rate=bad_amount / amount * 100, rejection_reasons={})
            for pid, name, source, campaign, content, amount, bad_amount in rows]
        reason = sa.func.left(sa.func.coalesce(sa.func.nullif(
            sa.func.split_part(lead.validation_reason, ':', 1), ''), 'unknown'), 200)
        reasons = db.execute(sa.select(reason, count).select_from(lead)
            .join(project, lead.project_id == project.id).where(*filters, rejected_condition())
            .group_by(reason).order_by(count.desc(), reason).limit(10)).all()
        report = SimpleNamespace(period_start=start, period_end=end, total_leads=total,
            total_rejected=rejected, overall_rejection_rate=rejected / total * 100 if total else 0,
            bad_sources=sources, top_rejection_reasons=dict(reasons),
            other_rejection_count=rejected - sum(n for _, n in reasons))
        return report, get_blacklist(db, owner_id)
    finally:
        db.rollback()  # No connection is held while openpyxl renders the file.


def spreadsheet_text(value):
    """Keep untrusted project/UTM names as text, never spreadsheet formulas."""
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def render_xlsx(report, blacklist):
    from io import BytesIO
    from openpyxl import Workbook
    book = Workbook()
    summary = book.active
    summary.title = 'Сводка'
    summary.append(['Всего', 'Отклонено', 'Начало UTC', 'Конец UTC (не включён)', 'Другие причины'])
    summary.append([report.total_leads, report.total_rejected, report.period_start.isoformat(),
                    report.period_end.isoformat(), report.other_rejection_count])
    if report.bad_sources:
        sheet = book.create_sheet('Плохие источники')
        sheet.append(['Проект', 'ID проекта', 'Источник', 'Кампания', 'Площадка',
                      'Всего заявок', 'Отклонено', 'Процент мусора'])
        for s in report.bad_sources:
            sheet.append([spreadsheet_text(v) for v in [s.project_name, s.project_id, s.source,
                s.campaign, s.content, s.total_leads, s.rejected_leads, round(s.rejection_rate, 2)]])
    if report.top_rejection_reasons:
        sheet = book.create_sheet('Причины отклонения')
        sheet.append(['Причина', 'Количество'])
        for reason, count in report.top_rejection_reasons.items():
            sheet.append([spreadsheet_text(reason), count])
    if blacklist:
        sheet = book.create_sheet('Чёрный список')
        columns = ['project_id', 'project_name', 'source', 'campaign', 'content', 'reason', 'ttl_seconds', 'expires_in_days']
        sheet.append(columns)
        for row in blacklist:
            sheet.append([spreadsheet_text(row[name]) for name in columns])
    for sheet in book:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
