from datetime import datetime, timedelta, timezone
from importlib import import_module, util
from pathlib import Path
from io import BytesIO
import uuid
import pytest
import sqlalchemy as sa
from alembic.migration import MigrationContext
from alembic.operations import Operations
from openpyxl import load_workbook
from core import models
from core.job_fence import fenced_job, LeaseLost
from automation import lead_placement_work as work, lead_alert_work as planner
from automation.work_tables import jobs
from automation.work_errors import RejectedBeforeExternalIO
from lead_validator.services import scoped_placements as blocks, quality_report
from tests.test_lead_scoped_stats import scope as base_scope, seed
from tests.test_durable_work import pg, claim, expire


@pytest.fixture
def scope(base_scope):
    base_scope.now = datetime.now(timezone.utc).replace(second=0, microsecond=0)
    return base_scope


def plan(s):
    s.now = datetime.now(timezone.utc).replace(second=0, microsecond=0)
    planner.plan_page(s.factory, 'lead.blacklist', {'scheduled_at': s.now.isoformat(), 'owner_id': str(s.owner)})
    with s.factory() as db:
        job = db.execute(sa.select(jobs).where(jobs.c.kind == 'lead.blacklist.project',
            jobs.c.tenant == str(s.owner))).mappings().one()
    s.job, s.payload = job['id'], job['payload']
    s.token = claim(s.factory, s.job)['lease_token']


def run(s):
    with fenced_job(s.job, s.token):
        return work.execute(s.factory, s.payload)


def bad(s, count=10, **overrides):
    for i in range(count):
        seed(s, utm_source='yandex', utm_campaign='campaign', utm_content='placement', **overrides)


def test_scoped_generation_and_repeat_do_not_extend_expiry(scope):
    s = scope
    bad(s); bad(s, project_id=s.foreign)
    plan(s)  # No Telegram chat is required for a blacklist job.
    assert run(s) == {'candidates': 1, 'active': 1}
    with s.factory() as db:
        expiry = db.scalar(sa.select(models.LeadPlacementBlock.expires_at))
        assert blocks.is_blacklisted(db, s.owner, s.project, 'yandex', 'campaign', 'placement')
        assert not blocks.is_blacklisted(db, s.other, s.foreign, 'yandex', 'campaign', 'placement')
        assert not blocks.is_blacklisted(db, s.owner, s.foreign, 'yandex', 'campaign', 'placement')
        assert len(blocks.get_blacklist(db, s.owner)) == 1
        assert blocks.get_blacklist(db, s.other) == []
    assert run(s) == {'candidates': 1, 'active': 1}
    with s.factory() as db:
        assert db.scalar(sa.select(models.LeadPlacementBlock.expires_at)) == expiry


def test_window_pending_and_self_rejections_do_not_trigger_blocks(scope):
    s = scope
    bad(s, status=models.LeadStatus.PENDING)
    bad(s, created_at=datetime.now(timezone.utc) - timedelta(days=22))
    bad(s, created_at=datetime.now(timezone.utc) + timedelta(days=1))
    bad(s, validation_reason='utm_invalid:blacklisted_placement:dynamic:yandex/campaign/placement')
    plan(s)
    assert run(s)['active'] == 0


def test_null_empty_and_colon_keys_do_not_cross_projects(scope):
    s = scope
    for i in range(10):
        seed(s, utm_source=None if i % 2 else '', utm_campaign=None, utm_content=None)
    plan(s); run(s)
    with s.factory() as db:
        assert blocks.is_blacklisted(db, s.owner, s.project, None, '', None)
    assert blocks.key(('a:b', 'c', 'd')) != blocks.key(('a', 'b:c', 'd'))


@pytest.mark.parametrize('change', ['owner', 'paused', 'scope', 'policy', 'lease'])
def test_stale_binding_cannot_write(scope, monkeypatch, change):
    s = scope
    bad(s); plan(s)
    if change == 'lease':
        expire(s.factory, s.job)
    elif change == 'policy':
        monkeypatch.setattr(work.settings, 'PLACEMENT_BLACKLIST_MIN_LEADS', 11)
    elif change == 'scope':
        s.payload = {**s.payload, 'scope_digest': 'forged'}
    else:
        with s.factory.begin() as db:
            p = db.get(models.PhoneProject, s.project)
            if change == 'owner':
                p.owner_id = s.other
            else:
                p.is_active = False
    with pytest.raises((RejectedBeforeExternalIO, LeaseLost)):
        run(s)
    with s.factory() as db:
        assert db.scalar(sa.select(sa.func.count()).select_from(models.LeadPlacementBlock)) == 0


def test_bounded_generation_rolls_back_instead_of_partial_blocklist(scope, monkeypatch):
    s = scope
    bad(s)
    for i in range(10):
        seed(s, utm_source='other')
    plan(s)
    monkeypatch.setattr(work, 'MAX_BLOCKS', 1)
    with pytest.raises(RejectedBeforeExternalIO, match='Too many'):
        run(s)
    with s.factory() as db:
        assert blocks.get_blacklist(db, s.owner) == []


def test_expired_or_transferred_decisions_are_not_applied(scope):
    s = scope
    bad(s); plan(s); run(s)
    with s.factory.begin() as db:
        db.get(models.PhoneProject, s.project).owner_id = s.other
    with s.factory() as db:
        assert blocks.get_blacklist(db, s.owner) == []
        assert not blocks.is_blacklisted(db, s.other, s.project, 'yandex', 'campaign', 'placement')
    with s.factory.begin() as db:
        db.get(models.PhoneProject, s.project).owner_id = s.owner
        db.execute(sa.update(models.LeadPlacementBlock).values(expires_at=datetime.now(timezone.utc) - timedelta(seconds=1)))
    with s.factory() as db:
        assert blocks.get_blacklist(db, s.owner) == []


@pytest.mark.asyncio
async def test_utm_uses_only_authorized_project_decisions(scope):
    from lead_validator.services.utm_validator import UTMValidator, UTMData
    s = scope
    bad(s); plan(s); run(s)
    validator = UTMValidator()
    validator.blacklisted_placements = []
    utm = UTMData(source='yandex', campaign='campaign', content='placement')
    with s.factory() as db:
        blocked = await validator.validate(utm, db=db, owner_id=s.owner, project_id=s.project)
        assert not blocked.is_valid and 'dynamic:' in blocked.reason
        assert (await validator.validate(utm, db=db, owner_id=s.other, project_id=s.foreign)).is_valid
    assert (await validator.validate(utm)).is_valid  # Unbound diagnostics never consult global state.
    validator.blacklisted_placements = ['placement']
    assert not (await validator.validate(utm)).is_valid  # Explicit static policy remains.


def test_quality_snapshot_honors_days_scope_and_releases_sql(scope):
    s = scope
    now = datetime.now(timezone.utc)
    bad(s, created_at=now - timedelta(hours=1)); bad(s, project_id=s.foreign)
    seed(s, created_at=now - timedelta(days=2), status=models.LeadStatus.VALID)
    seed(s, status=models.LeadStatus.PENDING)
    with s.factory() as db:
        report, blacklist = quality_report.snapshot(db, s.owner, 1)
        assert report.total_leads == 10 and report.total_rejected == 10
        assert len(report.bad_sources) == 1 and report.bad_sources[0].project_id == str(s.project)
        assert blacklist == [] and s.engine.pool.checkedout() == 0
        report, _ = quality_report.snapshot(db, s.owner, 7)
        assert report.total_leads == 11
        assert report.top_rejection_reasons == {'captcha_failed': 10}


def test_excel_empty_and_formula_strings_render_without_sql(scope, monkeypatch):
    routes = import_module('lead_validator.router')
    s = scope
    original = quality_report.render_xlsx
    def writer(*args, **kwargs):
        assert s.engine.pool.checkedout() == 0
        return original(*args, **kwargs)
    monkeypatch.setattr(quality_report, 'render_xlsx', writer)
    from types import SimpleNamespace
    with s.factory() as db:
        result = routes.get_quality_report(days=7, format='excel', current_user=SimpleNamespace(id=s.owner), db=db)
    assert result.status_code == 200
    assert load_workbook(BytesIO(result.body)).sheetnames == ['Сводка']
    for i in range(5):
        seed(s, utm_source='=1+1', utm_campaign='+formula')
    with s.factory() as db:
        result = routes.get_quality_report(days=7, format='excel', current_user=SimpleNamespace(id=s.owner), db=db)
    book = load_workbook(BytesIO(result.body))
    cells = list(book['Плохие источники'].iter_rows(min_row=2))[0]
    assert cells[2].value == "'=1+1" and cells[2].data_type == 's'
    assert cells[3].value == "'+formula" and cells[3].data_type == 's'


def test_migration_and_non_destructive_downgrade(scope):
    s = scope
    path = Path(__file__).resolve().parents[1] / 'alembic/versions/f46f708192a3_scoped_lead_placements.py'
    spec = util.spec_from_file_location('placement_migration', path)
    migration = util.module_from_spec(spec); spec.loader.exec_module(migration)
    with s.engine.begin() as conn:
        models.LeadPlacementBlock.__table__.drop(conn)
        with Operations.context(MigrationContext.configure(conn)):
            migration.upgrade()
            migration.downgrade()
        assert sa.inspect(conn).has_table('lead_placement_blocks')
        assert sa.inspect(conn).get_indexes('lead_placement_blocks')[0]['name'] == 'ix_lead_placement_owner_expiry'


def test_blacklist_calendar_occurrence_is_replay_safe(monkeypatch):
    from automation.work_control import occurrences
    monkeypatch.setenv('LEAD_ALERT_TIMEZONE', 'UTC')
    now = datetime(2026, 9, 22, 9, tzinfo=timezone.utc)
    assert ('lead.blacklist', 'maintenance', True, now.isoformat()) in list(occurrences(now, now))


def test_planner_is_bounded_and_includes_projects_without_chat(scope):
    s = scope
    with s.factory.begin() as db:
        db.delete(db.get(models.PhoneProject, s.project))
        for i in range(205):
            db.add(models.PhoneProject(id=uuid.UUID(int=i + 1), owner_id=s.owner, name='Synthetic'))
    payload = {'scheduled_at': s.now.isoformat(), 'owner_id': str(s.owner)}
    for cursor, expected in [(None, 100), (100, 100), (200, 5)]:
        page = {**payload}
        if cursor:
            page.update(cursor=str(uuid.UUID(int=cursor)), upper_id=str(uuid.UUID(int=205)))
        result = planner.plan_page(s.factory, 'lead.blacklist', page)
        assert result == {'planned': expected, 'has_next': cursor != 200}
        assert planner.plan_page(s.factory, 'lead.blacklist', page) == result
    with s.factory() as db:
        rows = db.execute(sa.select(jobs).where(jobs.c.kind == 'lead.blacklist.project')).mappings().all()
        assert len(rows) == 205 and all(row['replay_safe'] for row in rows)
        assert all(row['tenant'] == str(s.owner) for row in rows)


def test_lease_loss_at_commit_rolls_back_generated_blocks(scope, monkeypatch):
    from core.job_fence import current_fence
    s = scope
    bad(s); plan(s)
    original = work.key
    def expire_before_insert(values):
        reset = current_fence.set(None)
        try:
            expire(s.factory, s.job)
        finally:
            current_fence.reset(reset)
        return original(values)
    monkeypatch.setattr(work, 'key', expire_before_insert)
    with pytest.raises(LeaseLost):
        run(s)
    with s.factory() as db:
        assert blocks.get_blacklist(db, s.owner) == []


def test_linked_client_owner_change_hides_existing_blocks(scope):
    s = scope
    bad(s); plan(s); run(s)
    with s.factory.begin() as db:
        client = models.Client(owner_id=s.other, name='Other owner')
        db.add(client); db.flush()
        db.get(models.PhoneProject, s.project).client_id = client.id
    with s.factory() as db:
        assert blocks.get_blacklist(db, s.owner) == []


def test_quality_http_auth_bounds_and_no_global_state(scope, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from core import security
    from core.database import get_db
    from types import SimpleNamespace
    routes = import_module('lead_validator.router')
    s = scope
    bad(s); bad(s, project_id=s.foreign); plan(s); run(s)
    def forbidden(*args, **kwargs):
        raise AssertionError('Global analytics must not be accessed')
    monkeypatch.setattr(routes.analytics_service, 'generate_weekly_report', forbidden)
    app = FastAPI(); app.include_router(routes.router)
    def session():
        with s.factory() as db:
            yield db
    app.dependency_overrides[get_db] = session
    with TestClient(app) as client:
        for path in ('/reports/quality', '/reports/blacklist'):
            assert client.get(path).status_code in (401, 403)
        app.dependency_overrides[security.get_current_user] = lambda: SimpleNamespace(id=s.owner)
        response = client.get('/reports/quality', params={'days': 7})
        assert response.status_code == 200
        assert response.json()['overall']['total_leads'] == 10
        assert len(response.json()['blacklisted_placements']) == 1
        assert client.get('/reports/blacklist').json()['count'] == 1
        assert client.get('/reports/quality', params={'days': 366}).status_code == 422
        assert client.get('/reports/quality', params={'format': 'bogus'}).status_code == 422
        app.dependency_overrides[security.get_current_user] = lambda: SimpleNamespace(id=s.other)
        assert client.get('/reports/blacklist').json()['count'] == 0
        monkeypatch.setattr(blocks, 'MAX_BLOCKS', 0)
        app.dependency_overrides[security.get_current_user] = lambda: SimpleNamespace(id=s.owner)
        assert client.get('/reports/blacklist').status_code == 422
        assert client.get('/reports/quality').status_code == 422
