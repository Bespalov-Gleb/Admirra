"""File safety, ownership, exports and chat contract; isolated SQLite, no provider calls."""
import io
import unittest
import uuid
import zipfile
from datetime import datetime, timezone
from types import SimpleNamespace
from unittest.mock import patch

from ai.assistant import files


class ExtractionTests(unittest.TestCase):
    def test_unicode_and_filename(self):
        self.assertEqual(files.extract('План 1590 ₽'.encode(), 'plan.md'), 'План 1590 ₽')
        self.assertEqual(files.safe_name('../../test.txt'), 'test.txt')
        self.assertEqual(files.extract('План'.encode('utf-16'), 'text.txt'), 'План')

    def test_invalid_files_and_limits(self):
        for data, name in [(b'', 'a.txt'), (b'hi', 'a.exe'), (b'\x00', 'a.txt'),
                           (b'hello', 'a.pdf'), (b'hello', 'a.docx'),
                           (b'x' * (files.MAX_CHARS + 1), 'a.md')]:
            with self.subTest(name=name), self.assertRaises((ValueError, zipfile.BadZipFile)):
                files.extract(data, name)

    def test_pdf_text_and_scan(self):
        import pymupdf
        doc = pymupdf.open()
        page = doc.new_page()
        with self.assertRaisesRegex(ValueError, 'OCR'):
            files.extract(doc.tobytes(), 'scan.pdf')
        page.insert_text((72, 72), 'Revenue 1590 RUB')
        self.assertIn('1590', files.extract(doc.tobytes(), 'report.pdf'))
        doc.close()

    def test_docx_and_export_markdown(self):
        from docx import Document
        text = '# Отчёт\n\n**Расход** 1 590 ₽\n- Проверить цели\n\n| Канал | Лиды |\n| --- | --- |\n| VK | 34 |\n\n```\ncode\n```'
        result = files.export_docx(text)
        doc = Document(io.BytesIO(result))
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(doc.tables[0].cell(1, 1).text, '34')
        self.assertIn('1 590 ₽', files.extract(result, 'report.docx'))
        self.assertTrue(any(r.bold for p in doc.paragraphs for r in p.runs))

    def test_zip_bomb_rejected(self):
        data = io.BytesIO()
        with zipfile.ZipFile(data, 'w', zipfile.ZIP_DEFLATED) as z:
            z.writestr('word/document.xml', b'a' * (21 * 1024 * 1024))
        with self.assertRaisesRegex(ValueError, 'распакованный'):
            files.extract(data.getvalue(), 'bad.docx')

    def test_dtd_rejected(self):
        data = io.BytesIO()
        with zipfile.ZipFile(data, 'w') as z:
            z.writestr('word/document.xml', '<!DOCTYPE x [<!ENTITY a "bad">]><x>&a;</x>')
        with self.assertRaisesRegex(ValueError, 'структура'):
            files.extract(data.getvalue(), 'bad.docx')


class IsolatedParserTests(unittest.IsolatedAsyncioTestCase):
    async def test_25mb_pdf_boundary(self):
        import pymupdf
        with pymupdf.open() as doc:
            doc.new_page().insert_text((72, 72), 'Boundary 34')
            data = doc.tobytes()
        data += b'\n' * (files.MAX_BYTES - len(data))
        self.assertIn('34', await files.extract_isolated(data, 'boundary.pdf'))
        with self.assertRaisesRegex(ValueError, '25'):
            await files.extract_isolated(data + b'\n', 'boundary.pdf')

    async def test_real_child_and_rejection(self):
        self.assertEqual(await files.extract_isolated(b'hello', 'a.txt'), 'hello')
        with self.assertRaises(ValueError):
            await files.extract_isolated(b'not a zip', 'a.docx')


class RouterTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from sqlalchemy import Column, DateTime, ForeignKey, JSON, String, Text, Uuid, create_engine
        from sqlalchemy.orm import declarative_base, relationship, sessionmaker
        from sqlalchemy.pool import StaticPool
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from ai.assistant import router as routes
        cls.routes = routes
        Base = declarative_base()

        class User(Base):
            __tablename__ = 'users'
            id = Column(Uuid, primary_key=True, default=uuid.uuid4)

        class Conversation(Base):
            __tablename__ = 'conversations'
            id = Column(Uuid, primary_key=True, default=uuid.uuid4)
            user_id = Column(Uuid, ForeignKey('users.id'))
            client_id = Column(Uuid, nullable=True)
            title = Column(String)
            model = Column(String)
            effort = Column(String)
            updated_at = Column(DateTime, default=lambda: datetime.now(timezone.utc))
            messages = relationship('Message', cascade='all, delete-orphan', order_by='Message.created_at')

        class Message(Base):
            __tablename__ = 'messages'
            id = Column(Uuid, primary_key=True, default=uuid.uuid4)
            conversation_id = Column(Uuid, ForeignKey('conversations.id'))
            role = Column(String)
            content = Column(Text)
            tool_calls = Column(JSON)
            created_at = Column(DateTime, default=lambda: datetime.now(timezone.utc))

        cls.engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
        cls.Base, cls.Session = Base, sessionmaker(bind=cls.engine)
        cls.models = SimpleNamespace(User=User, AiConversation=Conversation, AiMessage=Message)
        cls.app = FastAPI()
        cls.app.include_router(routes.router)
        cls.client = TestClient(cls.app)

    def setUp(self):
        self.Base.metadata.create_all(self.engine)
        with self.Session() as db:
            u = self.models.User()
            other = self.models.User()
            db.add_all([u, other]); db.flush()
            self.uid, self.other = u.id, other.id
            conv = self.models.AiConversation(user_id=u.id)
            foreign = self.models.AiConversation(user_id=other.id)
            db.add_all([conv, foreign]); db.flush()
            self.cid, self.foreign = str(conv.id), str(foreign.id)
            answer = self.models.AiMessage(conversation_id=conv.id, role='assistant', content='# Ответ\n**34** лида')
            db.add(answer); db.flush(); self.mid = str(answer.id)
            db.commit()

        def get_db():
            with self.Session() as db:
                yield db

        self.app.dependency_overrides[self.routes.get_db] = get_db
        self.app.dependency_overrides[self.routes.security.get_current_user] = lambda: SimpleNamespace(id=self.uid)
        self.patcher = patch.object(self.routes, 'models', self.models)
        self.patcher.start()

    def tearDown(self):
        self.patcher.stop()
        self.app.dependency_overrides.clear()
        self.Base.metadata.drop_all(self.engine)

    def upload(self, cid=None, text=b'hello'):
        return self.client.post(f'/assistant/conversations/{cid or self.cid}/attachments?filename=test.txt', content=text,
                                headers={'Content-Type': 'application/octet-stream'})

    def test_upload_pending_delete_and_privacy(self):
        uploaded = self.upload()
        self.assertEqual(uploaded.status_code, 201, uploaded.text)
        aid = uploaded.json()['id']
        self.assertNotIn('content', uploaded.json())
        conv = self.client.get(f'/assistant/conversations/{self.cid}').json()
        self.assertEqual(len(conv['pending_attachments']), 1)
        self.assertEqual(len(conv['messages']), 1)
        self.assertEqual(self.upload(self.foreign).status_code, 404)
        self.assertEqual(self.client.delete(f'/assistant/conversations/{self.foreign}/attachments/{aid}').status_code, 404)
        self.assertEqual(self.client.delete(f'/assistant/conversations/{self.cid}/attachments/{aid}').status_code, 204)

    def test_export_owner_only(self):
        for fmt in ('md', 'docx'):
            url = f'/assistant/conversations/{self.cid}/messages/{self.mid}/download?format={fmt}'
            response = self.client.get(url)
            self.assertEqual(response.status_code, 200)
            self.assertIn('attachment;', response.headers['content-disposition'])
            self.assertEqual(response.headers['cache-control'], 'private, no-store')
            if fmt == 'md': self.assertIn('**34**', response.text)
            else: self.assertTrue(zipfile.is_zipfile(io.BytesIO(response.content)))
        self.assertEqual(self.client.get(f'/assistant/conversations/{self.foreign}/messages/{self.mid}/download?format=md').status_code, 404)

    def test_chat_missing_attachment_and_count(self):
        response = self.client.post('/assistant/chat', json={'message': 'hi', 'conversation_id': self.cid,
                                                            'attachment_ids': [str(uuid.uuid4())]})
        self.assertEqual(response.status_code, 404)
        response = self.client.post('/assistant/chat', json={'message': 'hi', 'conversation_id': self.cid,
                                                            'attachment_ids': [str(uuid.uuid4()) for _ in range(4)]})
        self.assertEqual(response.status_code, 422)

    def test_history_keeps_document_once_and_hides_unused(self):
        from ai.assistant.agent import _history_to_messages
        a = SimpleNamespace(id=uuid.uuid4(), role='attachment', content='source 34', tool_calls={'filename': 'report.txt'})
        msg = SimpleNamespace(role='user', content='analyze', tool_calls={'attachments': [files.attachment_public(a)]})
        unused = SimpleNamespace(id=uuid.uuid4(), role='attachment', content='UNUSED', tool_calls={})
        history = _history_to_messages(SimpleNamespace(messages=[a, unused, msg, msg]))
        self.assertEqual(len(history), 2)
        self.assertIn('source 34', history[0]['content'])
        self.assertNotIn('source 34', history[1]['content'])
        self.assertNotIn('UNUSED', str(history))

    def test_delete_conversation_removes_attachments(self):
        aid = self.upload().json()['id']
        self.assertEqual(self.client.delete(f'/assistant/conversations/{self.cid}').status_code, 204)
        with self.Session() as db:
            self.assertIsNone(db.get(self.models.AiMessage, uuid.UUID(aid)))

    def test_used_attachment_cannot_be_removed(self):
        aid = self.upload().json()['id']
        with self.Session() as db:
            db.add(self.models.AiMessage(conversation_id=uuid.UUID(self.cid), role='user', content='hi',
                                         tool_calls={'attachments': [{'id': aid, 'name': 'test.txt'}]}))
            db.commit()
        self.assertEqual(self.client.delete(f'/assistant/conversations/{self.cid}/attachments/{aid}').status_code, 409)
        data = self.client.get(f'/assistant/conversations/{self.cid}').json()
        self.assertEqual(data['pending_attachments'], [])
        self.assertEqual(data['messages'][-1]['attachments'][0]['id'], aid)

    def test_conversation_text_quota(self):
        for _ in range(3):
            self.assertEqual(self.upload(text=b'a' * files.MAX_CHARS).status_code, 201)
        self.assertEqual(self.upload().status_code, 409)

    def test_raw_size_limit(self):
        self.assertEqual(self.upload(text=b'a' * (files.MAX_BYTES + 1)).status_code, 413)

    def test_chat_passes_attachment_and_done_id(self):
        aid = self.upload().json()['id']
        seen = []
        async def run(db, conv, text, model, effort, user, attachments=None):
            seen.extend(attachments)
            yield {'type': 'done', 'content': 'Готово', 'message_id': self.mid}
        with patch.object(self.routes.agent, 'run', run):
            response = self.client.post('/assistant/chat', json={
                'message': 'Прочитай', 'conversation_id': self.cid, 'attachment_ids': [aid]})
        self.assertEqual(response.status_code, 200)
        self.assertIn(self.mid, response.text)
        self.assertEqual(len(seen), 1)

    def test_unauthed_upload(self):
        from fastapi import HTTPException
        def no_user(): raise HTTPException(401, 'Unauthorized')
        self.app.dependency_overrides[self.routes.security.get_current_user] = no_user
        self.assertEqual(self.upload().status_code, 401)


if __name__ == '__main__':
    unittest.main()
